import os
import io
import pandas as pd
import numpy as np
import streamlit as st
import matplotlib.pyplot as plt
import seaborn as sns

from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import RandomForestClassifier
from sklearn.cluster import KMeans
from sklearn.metrics import accuracy_score, roc_auc_score, precision_score, recall_score, f1_score

from imblearn.over_sampling import SMOTE
from docx import Document

import google.generativeai as genai

# ================= CONFIG =================
st.set_page_config(page_title="Hệ thống Cảnh báo Tiểu đường", page_icon="🏥", layout="wide")

# Google Gemini API key: store this in Streamlit Secrets, NOT in GitHub.
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
except Exception:
    GOOGLE_API_KEY = ""

if GOOGLE_API_KEY:
    genai.configure(api_key=GOOGLE_API_KEY)

# ================= UI & CSS (BẢNG MÀU MỚI - ĐỘ TƯƠNG PHẢN CAO) =================
st.markdown("""
<style>
    /* Nền xanh nước biển nhạt làm nổi bật các khối màu trắng */
    .stApp, [data-testid="stAppViewContainer"] {
        background-color: #EFF7F5;
        color: #163B45;
    }

    /* Giữ chữ dễ đọc khi Streamlit đang dùng theme tối */
    .stApp [data-testid="stMarkdownContainer"],
    .stApp [data-testid="stMarkdownContainer"] p,
    .stApp [data-testid="stMarkdownContainer"] li,
    .stApp [data-testid="stWidgetLabel"],
    .stApp [data-testid="stWidgetLabel"] p,
    .stApp label {
        color: #52666B !important;
    }
    
    /* Container chính */
    .block-container { 
        max-width: 1100px; 
        padding-top: 2rem; 
        padding-bottom: 2rem;
    }
    
    /* Box tiêu đề chính */
    .medical-header {
        color: #163B45 !important; /* Xanh than */
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
        font-weight: 800;
        text-align: center;
        padding: 20px;
        background-color: #FFFFFF;
        border-radius: 12px;
        box-shadow: 0px 4px 10px rgba(0,0,0,0.05);
        margin-bottom: 30px;
        border-bottom: 5px solid #168A83;
    }
    
    /* Form nhập liệu (Card trắng nổi lên) */
    div[data-testid="stForm"] {
        background-color: #FFFFFF;
        padding: 25px;
        border-radius: 12px;
        box-shadow: 0 4px 15px rgba(0,0,0,0.06);
        border: 1px solid #D8E7E3;
    }
    
    /* Style cho các thẻ hiển thị kết quả */
    div[data-testid="metric-container"] {
        background-color: #FFFFFF;
        border-left: 6px solid #168A83;
        padding: 15px 20px;
        border-radius: 8px;
        box-shadow: 0 3px 10px rgba(0,0,0,0.08);
    }

    /* Ép màu chữ cho Metric để không bị lỗi Darkmode */
    div[data-testid="metric-container"] label {
        color: #52666B !important;
        font-weight: 600;
    }

    /* Ô nhập số có nền sáng và chữ tối */
    div[data-testid="stNumberInput"] input {
        background-color: #FFFFFF !important;
        color: #163B45 !important;
        border-color: #D8E7E3 !important;
    }
    div[data-testid="stNumberInput"] button {
        background-color: #FFFFFF !important;
        color: #52666B !important;
    }
    
    /* Tiêu đề các phần */
    .section-title {
        color: #163B45 !important;
        font-size: 1.4rem;
        font-weight: 700;
        margin-top: 25px;
        margin-bottom: 15px;
        border-left: 5px solid #168A83;
        padding-left: 12px;
    }

    /* Nút Submit chính */
    div[data-testid="stFormSubmitButton"]>button {
        background-color: #168A83;
        color: white !important;
        font-size: 1.1rem;
        font-weight: 700;
        border-radius: 8px;
        border: none;
        height: 3.2em;
        width: 100%;
        transition: all 0.3s ease;
    }
    div[data-testid="stFormSubmitButton"]>button:hover {
        background-color: #0F6E69;
        box-shadow: 0 4px 12px rgba(15, 76, 117, 0.3);
    }

    /* Các nút phụ (Tải file, Tìm phòng khám) */
    .stDownloadButton>button, .stLinkButton>a {
        width: 100%;
        border-radius: 8px;
        font-weight: 600;
        height: 3em;
        background-color: #FFFFFF;
        color: #163B45 !important;
        border: 1.5px solid #168A83;
        transition: 0.3s;
    }
    .stDownloadButton>button:hover, .stLinkButton>a:hover {
        background-color: #168A83;
        color: #FFFFFF !important;
    }

    /* Ép màu chữ cho các Tab */
    .stApp .stTabs [data-baseweb="tab-list"] button {
        color: #52666B !important;
        font-weight: 600;
        font-size: 1.1rem;
    }
    .stApp .stTabs [data-baseweb="tab-list"] button[aria-selected="true"] {
        color: #163B45 !important;
        border-bottom-color: #168A83 !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown('<h1 class="medical-header">🏥 CỔNG THÔNG TIN ĐÁNH GIÁ NGUY CƠ TIỂU ĐƯỜNG</h1>', unsafe_allow_html=True)

# ================= STATE MANAGEMENT =================
if 'prediction_done' not in st.session_state:
    st.session_state.prediction_done = False
if 'advice' not in st.session_state:
    st.session_state.advice = ""

# ================= LOAD DATA =================
@st.cache_data
def load_data():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    csv_path = os.path.join(base_dir, "diabetes.csv")
    return pd.read_csv(csv_path)

try:
    df = load_data()
except FileNotFoundError:
    st.error("⚠️ Không tìm thấy file 'diabetes.csv'. Vui lòng kiểm tra lại thư mục.")
    st.stop()

# ================= TRAIN =================
@st.cache_resource
def train(df):
    X = df.drop("Outcome", axis=1)
    y = df["Outcome"]

    cols = ['Glucose','BloodPressure','SkinThickness','Insulin','BMI']
    X[cols] = SimpleImputer(strategy='median').fit_transform(X[cols])

    scaler = StandardScaler()
    X_scaled = scaler.fit_transform(X)

    X_train, X_test, y_train, y_test = train_test_split(
        X_scaled, y, stratify=y, test_size=0.2, random_state=42
    )

    y_train_counts = y_train.value_counts()
    X_res, y_res = SMOTE(random_state=42).fit_resample(X_train, y_train)
    y_res_counts = y_res.value_counts()

    rf = RandomForestClassifier(n_estimators=200, random_state=42)
    rf.fit(X_res, y_res)

    lr = LogisticRegression(max_iter=1000)
    lr.fit(X_res, y_res)

    kmeans = KMeans(n_clusters=3, n_init=10, random_state=42)
    clusters = kmeans.fit_predict(X_scaled)

    return rf, lr, scaler, kmeans, X_test, y_test, clusters, X.columns, y_train_counts, y_res_counts

rf, lr, scaler, kmeans, X_test, y_test, clusters, feature_names, y_train_counts, y_res_counts = train(df)
df["Cụm"] = clusters

# ================= LABEL CLUSTER =================
cluster_mean = df.groupby("Cụm")[["Glucose","BMI"]].mean()

def label_cluster(row):
    if row["Glucose"] > 140:
        return "Nguy cơ cao"
    elif row["Glucose"] > 120:
        return "Tiền tiểu đường"
    else:
        return "Bình thường"

mapping = cluster_mean.apply(label_cluster, axis=1).to_dict()
df["Nhóm"] = df["Cụm"].map(mapping)

# ================= AI AGENT =================
def ai_agent(data, prob, cluster_label):
    if not GOOGLE_API_KEY:
        return "⚠️ Chưa cấu hình API Key cho AI."

    model = genai.GenerativeModel("gemini-2.5-flash")
    risk = "Nguy cơ cao" if prob > 0.6 else "Nguy cơ trung bình" if prob > 0.3 else "Nguy cơ thấp"

    prompt = f"""
    Bạn là chuyên gia tư vấn dinh dưỡng và chăm sóc sức khỏe.
    Thông tin bệnh nhân: Tuổi: {data['Age']}, BMI: {data['BMI']}, Glucose: {data['Glucose']}, Huyết áp: {data['BloodPressure']}
    Kết quả từ mô hình Học máy: Xác suất mắc tiểu đường: {prob*100:.1f}%, Mức: {risk}, Nhóm: {cluster_label}
    
    YÊU CẦU BẮT BUỘC:
    1. TUYỆT ĐỐI KHÔNG chẩn đoán lại hay dự đoán bệnh (việc này đã do mô hình học máy thực hiện).
    2. CHỈ sử dụng thông tin trên để đưa ra các gợi ý thiết thực về:
       - Chế độ ăn uống (thực phẩm nên ăn/tránh).
       - Chế độ tập luyện và sinh hoạt.
    Trình bày dưới 150 từ, sử dụng bullet points rõ ràng, giọng văn đồng cảm và chuyên nghiệp.
    """
    try:
        return model.generate_content(prompt).text
    except Exception as e:
        return f"Lỗi kết nối AI: {e}"

# ================= PREDICT =================
def predict(data):
    df_in = pd.DataFrame([data], columns=feature_names)
    x = scaler.transform(df_in)

    prob = rf.predict_proba(x)[0,1]
    pred = int(prob >= 0.4)
    cluster = kmeans.predict(x)[0]

    return prob, pred, cluster

# ================= DOCX GENERATOR =================
def create_report(data, prob, advice):
    doc = Document()
    doc.add_heading("BÁO CÁO KẾT QUẢ ĐÁNH GIÁ SỨC KHỎE", 0)
    doc.add_heading("1. Chỉ số lâm sàng đầu vào", level=1)
    for k, v in data.items():
        doc.add_paragraph(f"- {k}: {v}")
    
    doc.add_heading("2. Kết quả phân tích từ Mô hình Học máy", level=1)
    doc.add_paragraph(f"Xác suất dự đoán nguy cơ: {prob*100:.1f}%")
    
    doc.add_heading("3. Gợi ý chăm sóc sức khỏe (Tư vấn bởi AI)", level=1)
    doc.add_paragraph(advice)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ================= TABS =================
tab1, tab2, tab3 = st.tabs(["🩺 Bảng điều khiển Y tế", "📊 Phân tích Mô hình", "📈 Khám phá Dữ liệu"])

# ================= TAB 1: PREDICTION =================
with tab1:
    st.markdown('<div class="section-title">Nhập chỉ số bệnh nhân</div>', unsafe_allow_html=True)
    
    with st.form("prediction_form"):
        col1, col2, col3, col4 = st.columns(4)
        
        with col1:
            preg = st.number_input("Số lần mang thai", 0, 20, 1)
            glu = st.number_input("Đường huyết (Glucose)", 0, 300, 120)
        with col2:
            bp = st.number_input("Huyết áp (Diastolic)", 0, 200, 70)
            skin = st.number_input("Độ dày nếp gấp da", 0, 100, 20)
        with col3:
            ins = st.number_input("Insulin", 0, 1000, 79)
            bmi = st.number_input("Chỉ số BMI", 0.0, 70.0, 25.0)
        with col4:
            dpf = st.number_input("Chỉ số phả hệ (DPF)", 0.0, 3.0, 0.5)
            age = st.number_input("Tuổi", 1, 120, 30)

        submit_btn = st.form_submit_button("🔍 Tiến hành Đánh giá")

    if submit_btn:
        data = {
            "Pregnancies": preg, "Glucose": glu, "BloodPressure": bp, 
            "SkinThickness": skin, "Insulin": ins, "BMI": bmi, 
            "DiabetesPedigreeFunction": dpf, "Age": age
        }
        
        prob, pred, cluster = predict(data)
        cluster_label = df["Nhóm"].iloc[cluster]
        
        st.session_state.data = data
        st.session_state.prob = prob
        st.session_state.cluster_label = cluster_label
        
        with st.spinner("Đang trích xuất phác đồ chăm sóc từ hệ thống AI..."):
            st.session_state.advice = ai_agent(data, prob, cluster_label)
            st.session_state.prediction_done = True

    if st.session_state.prediction_done:
        st.markdown('<div class="section-title">Kết quả Đánh giá Lâm sàng</div>', unsafe_allow_html=True)
        
        res_col1, res_col2, res_col3 = st.columns(3)
        prob_val = st.session_state.prob
        
        with res_col1:
            st.metric("Tỉ lệ rủi ro", f"{prob_val*100:.1f}%")
        with res_col2:
            if prob_val > 0.6:
                st.error("🔴 CẢNH BÁO: NGUY CƠ CAO")
            elif prob_val > 0.3:
                st.warning("🟡 LƯU Ý: NGUY CƠ TRUNG BÌNH")
            else:
                st.success("🟢 AN TOÀN: NGUY CƠ THẤP")
        with res_col3:
            st.info(f"📍 Phân nhóm: {st.session_state.cluster_label}")

        st.markdown('<div class="section-title">Gợi ý Chăm sóc sức khỏe (AI)</div>', unsafe_allow_html=True)
        st.info(st.session_state.advice)

        st.markdown('<div class="section-title">Hành động & Hỗ trợ</div>', unsafe_allow_html=True)
        action_col1, action_col2 = st.columns(2)
        
        with action_col1:
            report_buffer = create_report(st.session_state.data, st.session_state.prob, st.session_state.advice)
            st.download_button(
                label="📄 Tải Báo Cáo Sức Khỏe (.docx)",
                data=report_buffer,
                file_name="Bao_Cao_Y_Te.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
            
        with action_col2:
            google_maps_query = "https://www.google.com/maps/search/phòng+khám+tiểu+đường+hoặc+bệnh+viện+gần+đây"
            st.link_button(
                "🏥 Tìm phòng khám gần tôi", 
                url=google_maps_query,
                use_container_width=True
            )

# ================= TAB 2 & 3 =================
with tab2:
    st.markdown('<div class="section-title">Đánh giá Hiệu năng Hệ thống Học máy</div>', unsafe_allow_html=True)
    models = {"Logistic Regression": lr, "Random Forest": rf}
    eval_data = []
    for name, model in models.items():
        y_pred = model.predict(X_test)
        y_prob = model.predict_proba(X_test)[:, 1]
        eval_data.append({
            "Mô hình": name,
            "Accuracy": accuracy_score(y_test, y_pred),
            "Precision": precision_score(y_test, y_pred),
            "Recall": recall_score(y_test, y_pred),
            "F1-Score": f1_score(y_test, y_pred),
            "ROC-AUC": roc_auc_score(y_test, y_prob)
        })

    df_eval = pd.DataFrame(eval_data)
    st.dataframe(df_eval.style.format({
        "Accuracy": "{:.2%}", "Precision": "{:.2%}", 
        "Recall": "{:.2%}", "F1-Score": "{:.2%}", "ROC-AUC": "{:.2%}"
    }), use_container_width=True)

    st.markdown('<div class="section-title">Trọng số Đặc trưng (Feature Importance)</div>', unsafe_allow_html=True)
    col_feat1, col_feat2 = st.columns([1, 2])
    imp = rf.feature_importances_
    feat_df = pd.DataFrame({"Feature": feature_names, "Importance": imp}).sort_values(by="Importance", ascending=False)
    
    with col_feat1:
        st.dataframe(feat_df, hide_index=True)
    with col_feat2:
        fig_feat, ax_feat = plt.subplots(figsize=(8, 4))
        sns.barplot(data=feat_df, x="Importance", y="Feature", color="#168A83", ax=ax_feat)
        st.pyplot(fig_feat)

with tab3:
    st.markdown('<div class="section-title">Xử lý Dữ liệu Mất cân bằng (SMOTE)</div>', unsafe_allow_html=True)
    col_smote1, col_smote2 = st.columns(2)
    with col_smote1:
        st.write("**Dữ liệu gốc (Train):**")
        st.bar_chart(y_train_counts, color="#E74C3C")
    with col_smote2:
        st.write("**Sau khi cân bằng (SMOTE):**")
        st.bar_chart(y_res_counts, color="#2ECC71")

    st.markdown("---")
    col_plot1, col_plot2 = st.columns(2)
    with col_plot1:
        st.markdown('<div class="section-title">Phân cụm Bệnh nhân (K-Means)</div>', unsafe_allow_html=True)
        fig_cluster, ax_cluster = plt.subplots(figsize=(6, 4))
        sns.scatterplot(data=df, x="Glucose", y="BMI", hue="Nhóm", palette="Set2", ax=ax_cluster)
        st.pyplot(fig_cluster)

    with col_plot2:
        st.markdown('<div class="section-title">Tương quan Chỉ số (Correlation)</div>', unsafe_allow_html=True)
        fig_corr, ax_corr = plt.subplots(figsize=(6, 4))
        sns.heatmap(df.select_dtypes(include=np.number).corr(), annot=False, cmap="Blues", ax=ax_corr, linewidths=.5)
        st.pyplot(fig_corr)